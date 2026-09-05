"""简历解析：支持 .pdf / .docx / .doc，提取姓名、教育、技能、经历、项目。"""
from __future__ import annotations

import io
import os
import re
from typing import List

from ..schemas import ResumeParseOut

PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
ALLOWED = PDF_EXT | DOCX_EXT | {".doc"}


def _read_pdf(raw: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n".join(parts).strip()


def _read_docx(raw: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(raw))
    lines = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines).strip()


def _read_doc_fallback(raw: bytes, filename: str) -> str:
    for enc in ("utf-8", "gbk", "latin-1"):
        try:
            s = raw.decode(enc, errors="ignore")
            s = re.sub(r"[\x00-\x08\x0B-\x1F\x7F]", "", s)
            s = re.sub(r"\s+", " ", s).strip()
            if len(s) > 50:
                return s
        except Exception:
            continue
    raise ValueError(f"无法解析文件 {filename}，请转换为 PDF/DOCX 后重新上传")


def parse_resume_bytes(raw: bytes, filename: str) -> ResumeParseOut:
    if not raw:
        raise ValueError("上传文件为空")
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in ALLOWED:
        raise ValueError(f"不支持的文件格式 {ext or '未知'}，仅支持 PDF / DOCX / DOC")

    if ext in PDF_EXT:
        text = _read_pdf(raw)
    elif ext in DOCX_EXT:
        text = _read_docx(raw)
    else:
        text = _read_doc_fallback(raw, filename)

    if len(text.strip()) < 20:
        raise ValueError("解析出的文本过短，文件可能损坏或为扫描件（需要可复制文字的 PDF/Word）")

    return ResumeParseOut(
        raw_text=text,
        name=_extract_name(text, filename),
        education=_extract_block(text, ["教育背景", "教育经历", "教育", "学历", "EDUCATION"]),
        skills=_extract_skills(text),
        experiences=_extract_block(text, ["工作经历", "工作经验", "职业经历", "实习经历", "EXPERIENCE", "WORK"]),
        projects=_extract_block(text, ["项目经验", "项目经历", "项目", "PROJECT"]),
    )


def _extract_name(text: str, filename: str) -> str:
    base = os.path.splitext(os.path.basename(filename or ""))[0]
    if re.fullmatch(r"[\u4e00-\u9fa5A-Za-z·\-]{2,20}", base or ""):
        return base
    for line in text.splitlines()[:10]:
        s = line.strip()
        m = re.search(r"姓\s*名[:：\s]*([\u4e00-\u9fa5A-Za-z·\-]{2,20})", s)
        if m:
            return m.group(1)
        if re.fullmatch(r"[\u4e00-\u9fa5]{2,4}", s):
            return s
    return ""


def _extract_block(text: str, keywords) -> List[str]:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    result, capture, current = [], False, ""
    for line in lines:
        hit = any((k in line.upper() if k.isascii() else k in line) for k in keywords)
        if hit and not capture:
            capture = True
            current = line
            continue
        if capture:
            if re.match(r"^[\u4e00-\u9fa5A-Z]{2,10}[\s:：]?$", line) and not any(
                (k in line.upper() if k.isascii() else k in line) for k in keywords
            ):
                if current.strip():
                    result.append(current.strip())
                capture, current = False, ""
                continue
            current = (current + "\n" + line).strip() if current else line
    if current.strip():
        result.append(current.strip())
    seen, out = set(), []
    for r in result:
        if 10 < len(r) < 800 and r not in seen:
            seen.add(r)
            out.append(r[:600])
        if len(out) >= 8:
            break
    return out


def _extract_skills(text: str) -> List[str]:
    TECH = [
        "Java", "Spring", "SpringBoot", "Spring Boot", "Spring Cloud", "MyBatis",
        "MySQL", "Oracle", "PostgreSQL", "Redis", "MongoDB", "Kafka", "RabbitMQ",
        "RocketMQ", "Docker", "Kubernetes", "K8s", "Nginx", "Linux", "Git",
        "Python", "Django", "Flask", "FastAPI", "Pandas", "TensorFlow", "PyTorch",
        "Go", "Golang", "Rust", "C++", "C#", "Node.js", "TypeScript", "JavaScript",
        "Vue", "React", "Angular", "HTML", "CSS", "Tailwind", "Webpack", "Vite",
        "JVM", "JUC", "分布式", "微服务", "高并发", "TCP/IP", "HTTP",
        "Hadoop", "Spark", "Flink", "Hive", "SQL", "ETL", "数据仓库", "机器学习",
        "深度学习", "NLP", "大模型", "LLM", "Prompt", "RAG", "LangChain",
    ]
    found = []
    for kw in TECH:
        if re.search(r"(?<![A-Za-z])" + re.escape(kw) + r"(?![A-Za-z])", text, flags=re.I):
            found.append(kw)
        if len(found) >= 20:
            break
    return found
